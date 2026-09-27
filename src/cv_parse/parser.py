"""Extract skills / years / level from PDF or DOCX CVs."""

from __future__ import annotations

import io
from pathlib import Path

from .experience import parse_experience
from .level import parse_level


def _read_pdf(data: bytes) -> str:
    import pdfplumber

    texts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
    return "\n".join(texts)


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return _read_pdf(data)
    if name.endswith(".docx"):
        return _read_docx(data)
    return data.decode("utf-8", errors="ignore")


def parse_cv(filename: str, data: bytes, extractor=None) -> dict:
    if extractor is None:
        from src.features.skills import SkillExtractor

        extractor = SkillExtractor()
    text = extract_text(filename, data)
    skills = extractor.extract(text)
    exp = parse_experience(text)
    level = parse_level(text)
    return {
        "text": text[:8000],
        "skills": skills,
        "years_mid": exp["years_mid"],
        "level_norm": level["level_norm"],
        "level_ordinal": level["level_ordinal"],
    }


def parse_cv_path(path: str | Path) -> dict:
    p = Path(path)
    return parse_cv(p.name, p.read_bytes())
